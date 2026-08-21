"""
Compare Alfred (generic LLM) vs ST GitHub Analyzer (KB-powered persona).

Sends the same set of technical STM32 questions to both systems, collects
responses side-by-side, and exports an Excel file for evaluation.

Goal: demonstrate the added value of a dedicated KB + pipeline approach
over a generic LLM that has no access to GitHub issues/documentation.

Usage:
    python pipeline_Automation/test_suites/compare_alfred_vs_persona.py
    python pipeline_Automation/test_suites/compare_alfred_vs_persona.py --max-questions 10
    python pipeline_Automation/test_suites/compare_alfred_vs_persona.py --use-proxy

Output:
    datasets/07_delivery/st_ready/evaluation_report/comparison_alfred_vs_persona.xlsx

Requires: openpyxl, requests
"""

import hashlib
import os
import random
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Optional
import argparse
import json

import requests
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    raise SystemExit(1)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = PROJECT_ROOT / "datasets" / "07_delivery" / "st_ready" / "evaluation_report"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ============================================================
# API CONFIGURATION
# ============================================================

# --- Alfred (generic st_copilot) ---
ALFRED_CLIENT_APP = os.getenv("ALFRED_CLIENT_APP_NAME", "mdrf_stgithub_analyzer_client").strip()
ALFRED_API_KEY = (
    os.getenv("ST_CHATGPT_API_KEY")
    or os.getenv("ST_AI_BRIDGE_API_KEY")
    or os.getenv("ST_API_KEY")
    or ""
).strip()
ALFRED_PERSONA = "st_copilot"

# --- ST GitHub Analyzer (KB-powered persona) ---
PERSONA_CLIENT_APP = os.getenv("PERSONA_CLIENT_APP_NAME", "mdrf_st_github_analyzer").strip()
PERSONA_API_KEY = (
    os.getenv("PERSONA_API_KEY")
    or os.getenv("ST_GITHUB_ANALYZER_API_KEY")
    or os.getenv("ST_CHATGPT_API_KEY")
    or os.getenv("ST_AI_BRIDGE_API_KEY")
    or os.getenv("ST_API_KEY")
    or ""
).strip()
PERSONA_NAME = "ST GitHub Analyzer"

# Shared
API_URL = "https://api-ai-bridge.st.com/chatgpt/api/client-apps"
CHAT_SERVICE = "chat"
PROXIES = {
    "http": "http://185.46.212.88:80",
    "https": "http://185.46.212.88:80",
}

# ============================================================
# EVALUATION QUESTIONS
# Designed to highlight differences: KB-specific vs generic knowledge
# ============================================================

COMPARISON_QUESTIONS = [
    # --- Questions where KB should clearly win (issue-specific) ---
    {
        "question": "My STM32H743 project crashes with a hard fault when I enable DMA for SPI. SPI works in polling mode. What could be wrong?",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "H7",
        "why": "KB contains real GitHub issues with DMA alignment and MPU config solutions",
    },
    {
        "question": "I migrated my STM32H747I-DISCO Ethernet project from CubeH7 v1.9 to v1.11 and now LwIP doesn't get an IP address. PHY link is up but no DHCP. Any idea?",
        "category": "Migration/Version",
        "expected_advantage": "Persona",
        "tags": "H7",
        "why": "KB has release notes + issues documenting ETH descriptor changes in v1.10",
    },
    {
        "question": "USB CDC on STM32H750B-DK works for 30s then PC says 'USB device not recognized'. Same code works on F4.",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "H7",
        "why": "KB has multiple issues about H7 USB VBUS sensing and clock recovery",
    },
    {
        "question": "HAL_UART_Receive_IT on STM32F446RE only gives me the first character then stops. I have to call it again in the callback. Is this expected?",
        "category": "HAL Behavior",
        "expected_advantage": "Persona",
        "tags": "F4",
        "why": "KB has issues explaining this is by design + recommended patterns",
    },
    {
        "question": "After a slave sends NACK, my I2C1 on STM32F401 gets permanently stuck with BUSY flag. Only chip reset clears it. How to recover?",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "F4",
        "why": "KB has the famous I2C BUSY bug workaround (toggle SCL lines)",
    },
    {
        "question": "After enabling TrustZone on STM32H573I-Discovery, I can't toggle LEDs from non-secure code. GPIO init returns HAL_OK but nothing happens.",
        "category": "Security/TrustZone",
        "expected_advantage": "Persona",
        "tags": "H5",
        "why": "KB has issues about SAU/GTZC GPIO pin attribution in TZ mode",
    },
    {
        "question": "SPI DMA transfers on STM32H563 randomly return corrupted data (1 in 50 transfers). Same code on F4 is rock solid.",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "H5",
        "why": "KB has H5 GPDMA linked-list config issues and workarounds",
    },
    {
        "question": "My STM32WL Nucleo sends LoRa join requests but never gets join accept. Gateway sees the request. DevEUI/AppKey confirmed. What blocks the join?",
        "category": "LoRaWAN",
        "expected_advantage": "Persona",
        "tags": "WL",
        "why": "KB has issues about WL RF switch config (FE_CTRL pins) causing RX failure",
    },
    {
        "question": "STM32WL55 outputs only +10dBm even though I configured +22dBm with HP PA. Radio config matches datasheet. What am I missing?",
        "category": "RF/Radio",
        "expected_advantage": "Persona",
        "tags": "WL",
        "why": "KB has issues about SMPS vs LDO regulator selection for HP PA",
    },
    {
        "question": "FreeRTOS crashes when I read SD card via SDMMC1 on STM32H7B3I-DK. Works in bare metal. Stack overflow detection doesn't trigger.",
        "category": "RTOS Integration",
        "expected_advantage": "Persona",
        "tags": "H7",
        "why": "KB has SDMMC DMA buffer alignment + MPU config issues",
    },
    # --- Questions where both might answer well (general STM32) ---
    {
        "question": "What is the difference between HAL and LL drivers in STM32Cube packages?",
        "category": "General Knowledge",
        "expected_advantage": "Tie",
        "tags": "General",
        "why": "Basic STM32 knowledge, both should know this",
    },
    {
        "question": "How does the STM32 clock tree work? What are HSI, HSE, PLL?",
        "category": "General Knowledge",
        "expected_advantage": "Tie",
        "tags": "General",
        "why": "Fundamental STM32 concept in any documentation",
    },
    # --- Questions where KB has documentation-specific answers ---
    {
        "question": "What BSP examples come with the STM32H573I-Discovery board package? I'm looking for audio and display demos.",
        "category": "Documentation",
        "expected_advantage": "Persona",
        "tags": "H5",
        "why": "KB has the actual file list and README from the BSP package",
    },
    {
        "question": "What sensors and connectivity are available on the B-U585I-IOT02A board?",
        "category": "Documentation",
        "expected_advantage": "Persona",
        "tags": "U5",
        "why": "KB has the actual board BSP documentation with sensor list",
    },
    {
        "question": "What breaking changes should I know when updating from STM32CubeF4 v1.26 to v1.28? My project uses USB, Ethernet, and FreeRTOS.",
        "category": "Release Notes",
        "expected_advantage": "Persona",
        "tags": "F4",
        "why": "KB has the actual release notes with change details",
    },
    # --- Diagnostic/troubleshooting questions ---
    {
        "question": "ADC with DMA in scan mode on STM32F429 gives the same value in all buffer positions. 4 channels configured but buffer fills with channel 0 data repeated.",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "F4",
        "why": "KB has ADC multi-channel DMA config issues with specific fixes",
    },
    {
        "question": "I set up LPDMA to transfer ADC results to SRAM4 in STOP2 mode on STM32U585 but the transfer never starts. Works in Run mode.",
        "category": "Low Power",
        "expected_advantage": "Persona",
        "tags": "U5",
        "why": "KB has LPDMA domain/clock retention issues in STOP modes",
    },
    {
        "question": "FDCAN on STM32H753 receives frames fine but some messages with specific IDs are dropped even though my filter should accept them.",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "H7",
        "why": "KB has FDCAN filter element size and RAM config issues",
    },
    {
        "question": "I need continuous ADC with DMA on STM32H5 using GPDMA. Examples show linked-list mode. Is there a simpler way or is linked-list mandatory?",
        "category": "Architecture",
        "expected_advantage": "Persona",
        "tags": "H5",
        "why": "KB has H5 GPDMA architectural explanations from issues/docs",
    },
    {
        "question": "My STM32H743 ADC readings jump by 200 LSBs even with stable DC input. Averaging doesn't help. What am I doing wrong?",
        "category": "Issue-specific",
        "expected_advantage": "Persona",
        "tags": "H7",
        "why": "KB has the H743 ADC errata and BOOST/linearity calibration fixes",
    },
]


# ============================================================
# API CALL FUNCTIONS
# ============================================================


def _generate_token(client_app: str, api_key: str, timestamp: int, nonce: int) -> str:
    """Compute the SHA1 auth token expected by the ST AI Bridge chat API.

    Binds client app name, service, API key, timestamp and nonce together so
    the server can validate the caller without the raw API key being
    replayable on its own. Called separately for Alfred and for the persona
    since each uses its own client app name / API key pair.
    """
    data_string = f"{client_app}_{CHAT_SERVICE}_{api_key}_{timestamp}_{nonce}"
    return hashlib.sha1(data_string.encode("utf-8")).hexdigest()


def _extract_text(response_json: dict) -> Optional[str]:
    """Extract the answer text from an AI Bridge response, trying known shapes.

    Different personas/API versions have been observed to return the answer
    under different keys ("completion", "text", "content", "message",
    "answer", or OpenAI-style "choices[0].message.content"), so each shape is
    tried in turn. Returns None if nothing usable is found.
    """
    if not isinstance(response_json, dict):
        return None

    completion = response_json.get("completion")
    if isinstance(completion, str) and completion.strip():
        return completion.strip()

    for key in ("text", "content", "message", "answer"):
        val = response_json.get(key)
        if isinstance(val, str) and val.strip():
            return val.strip()

    choices = response_json.get("choices", [])
    if choices and isinstance(choices[0], dict):
        msg = choices[0].get("message", {})
        if isinstance(msg, dict) and msg.get("content"):
            return msg["content"].strip()

    return None


def ask_alfred(question: str, use_proxy: bool = False, retries: int = 3) -> str:
    """Send question to Alfred (generic st_copilot persona) and return its answer.

    Retries with exponential backoff (2, 4, 8... seconds) on HTTP errors or
    exceptions, since the AI Bridge endpoint can be transiently unavailable.
    A new timestamp/nonce/token is generated per attempt because the token
    embeds the timestamp and would be rejected as stale otherwise.
    """
    system_prompt = (
        "You are a technical support assistant for STM32 microcontrollers. "
        "Answer user questions about STM32Cube firmware, HAL drivers, BSP, "
        "middleware, and development boards. Be concise and technical."
    )

    for attempt in range(1, retries + 1):
        timestamp = int(time.time())
        nonce = random.randint(0, 999999)
        token = _generate_token(ALFRED_CLIENT_APP, ALFRED_API_KEY, timestamp, nonce)

        headers = {
            "Content-Type": "application/json",
            "stchatgpt-auth-token": token,
            "stchatgpt-auth-nonce": str(nonce),
        }

        payload = {
            "version": 1,
            "clientAppName": ALFRED_CLIENT_APP,
            "service": CHAT_SERVICE,
            "timestamp": timestamp,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": [{"type": "text", "content": question}]},
            ],
            "temperature": 0.3,
            "persona": ALFRED_PERSONA,
            "maxResponseTokens": 2048,
            "responseFormat": "text",
        }

        try:
            resp = requests.post(
                API_URL, json=payload, headers=headers,
                verify=False, timeout=120,
                proxies=PROXIES if use_proxy else None,
            )
            if resp.status_code >= 400:
                print(f"    [Alfred attempt {attempt}] HTTP {resp.status_code}")
                if attempt < retries:
                    time.sleep(2 ** attempt)
                    continue
                return f"(HTTP Error {resp.status_code})"

            data = resp.json()
            text = _extract_text(data)
            return text or f"(Unexpected response: {list(data.keys())})"

        except Exception as exc:
            print(f"    [Alfred attempt {attempt}] Error: {exc}")
            if attempt < retries:
                time.sleep(2 ** attempt)
                continue
            return f"(Error: {exc})"

    return "(No answer after retries)"


def ask_persona(question: str, use_proxy: bool = False, retries: int = 3) -> str:
    """Send question to ST GitHub Analyzer (KB-powered persona) and return its answer.

    Same retry/backoff strategy as ask_alfred, but with no system prompt and a
    plain-string user message: the persona's own system prompt and KB grounding
    are configured server-side, so this client only forwards the raw question.
    """
    for attempt in range(1, retries + 1):
        timestamp = int(time.time())
        nonce = random.randint(0, 999999)
        token = _generate_token(PERSONA_CLIENT_APP, PERSONA_API_KEY, timestamp, nonce)

        headers = {
            "Content-Type": "application/json",
            "stchatgpt-auth-token": token,
            "stchatgpt-auth-nonce": str(nonce),
        }

        payload = {
            "version": 1,
            "clientAppName": PERSONA_CLIENT_APP,
            "service": CHAT_SERVICE,
            "timestamp": timestamp,
            "messages": [
                {"role": "user", "content": question},
            ],
            "temperature": 0.3,
            "persona": PERSONA_NAME,
            "maxResponseTokens": 2048,
            "responseFormat": "text",
        }

        try:
            resp = requests.post(
                API_URL, json=payload, headers=headers,
                verify=False, timeout=120,
                proxies=PROXIES if use_proxy else None,
            )
            if resp.status_code >= 400:
                print(f"    [Persona attempt {attempt}] HTTP {resp.status_code}")
                if attempt < retries:
                    time.sleep(2 ** attempt)
                    continue
                return f"(HTTP Error {resp.status_code})"

            data = resp.json()
            text = _extract_text(data)
            return text or f"(Unexpected response: {list(data.keys())})"

        except Exception as exc:
            print(f"    [Persona attempt {attempt}] Error: {exc}")
            if attempt < retries:
                time.sleep(2 ** attempt)
                continue
            return f"(Error: {exc})"

    return "(No answer after retries)"


# ============================================================
# EVALUATION CRITERIA
# ============================================================

EVALUATION_COLUMNS = [
    "Cites specific GitHub issue/PR",
    "Mentions exact HAL function or errata",
    "Provides actionable fix (not just generic advice)",
    "References version-specific info",
    "Gives STM32-specific workaround",
]


# ============================================================
# EXCEL EXPORT
# ============================================================


def export_comparison_excel(results: list[dict], output_path: Path):
    """Export side-by-side Alfred vs Persona comparison results to a 3-sheet Excel workbook.

    Sheet 1 ("Comparison") holds one row per question with both answers plus
    empty columns for a human grader to fill in. Sheet 2 ("Grille d'évaluation")
    documents the scoring rubric. Sheet 3 ("Metadata") records run parameters
    for traceability. Called both incrementally (every 5 questions, as a
    crash-resilience checkpoint) and once at the very end.
    """
    wb = Workbook()

    # --- Sheet 1: Side-by-side comparison ---
    ws = wb.active
    ws.title = "Comparison"

    # Styles
    header_fill = PatternFill(start_color="1A2382", end_color="1A2382", fill_type="solid")
    header_font = Font(bold=True, size=10, color="FFFFFF")
    alfred_fill = PatternFill(start_color="FFE6CC", end_color="FFE6CC", fill_type="solid")
    persona_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Headers
    headers = [
        "N°", "Question", "Catégorie", "Tags",
        "Réponse Alfred (générique)", "Réponse ST GitHub Analyzer (KB)",
        "Avantage attendu", "Avantage réel",
        "Alfred cite issue?", "Persona cite issue?",
        "Alfred actionnable?", "Persona actionnable?",
        "Commentaire évaluation",
    ]
    col_widths = [4, 60, 18, 8, 80, 80, 14, 14, 14, 14, 14, 14, 40]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data rows
    for idx, result in enumerate(results, start=1):
        row_idx = idx + 1
        row_data = [
            idx,
            result["question"],
            result["category"],
            result["tags"],
            result["alfred_answer"],
            result["persona_answer"],
            result["expected_advantage"],
            "",  # To be filled manually
            "",  # Alfred cites issue?
            "",  # Persona cites issue?
            "",  # Alfred actionnable?
            "",  # Persona actionnable?
            "",  # Commentaire
        ]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align
            cell.border = thin_border
            cell.font = Font(size=9)
            # Color-code the two answer columns so graders can visually tell
            # which side (generic Alfred vs KB-powered persona) they are reading.
            if col_idx == 5:
                cell.fill = alfred_fill
            elif col_idx == 6:
                cell.fill = persona_fill

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"

    # --- Sheet 2: Scoring summary ---
    ws2 = wb.create_sheet("Grille d'évaluation")
    summary_headers = ["Critère", "Description", "Points si oui"]
    summary_data = [
        ["Cite une issue GitHub", "La réponse mentionne un numéro d'issue (#xxx) ou un lien GitHub", "2"],
        ["Mentionne une fonction HAL exacte", "Ex: HAL_SPI_Transmit_DMA, pas juste 'the SPI function'", "1"],
        ["Fix actionnable", "L'utilisateur peut appliquer la solution sans recherche supplémentaire", "2"],
        ["Info version-spécifique", "Mentionne une version de Cube/HAL où le bug existe ou est fixé", "2"],
        ["Workaround STM32-specific", "Solution propre au hardware ST (MPU config, errata, registre)", "2"],
        ["Source vérifiable", "Pointe vers un fichier, une release note, ou un commit précis", "1"],
    ]

    for col_idx, h in enumerate(summary_headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for row_idx, row_data in enumerate(summary_data, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = wrap_align

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 70
    ws2.column_dimensions["C"].width = 15

    # --- Sheet 3: Metadata ---
    ws3 = wb.create_sheet("Metadata")
    meta = [
        ["Date de génération", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Nombre de questions", str(len(results))],
        ["Alfred persona", ALFRED_PERSONA],
        ["ST GitHub Analyzer persona", PERSONA_NAME],
        ["API endpoint", API_URL],
        ["Temperature", "0.3"],
        ["Max response tokens", "2048"],
        ["Objectif", "Démontrer la valeur ajoutée de la KB GitHub vs LLM générique"],
    ]
    for row_idx, (key, val) in enumerate(meta, start=1):
        ws3.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        ws3.cell(row=row_idx, column=2, value=val)
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 80

    wb.save(output_path)
    print(f"\nExcel saved: {output_path}")


# ============================================================
# MAIN
# ============================================================


def main():
    """CLI entry point: ask both Alfred and the persona every question, then export.

    Validates that both sets of API credentials are present, resolves the
    question list (optionally truncated via --max-questions), calls both
    APIs sequentially per question (with a configurable delay to respect
    rate limits), checkpoints the Excel export every 5 questions, and
    finally writes both the Excel workbook and a raw JSON dump of the same
    results for further/automated analysis.
    """
    parser = argparse.ArgumentParser(
        description="Compare Alfred vs ST GitHub Analyzer responses"
    )
    parser.add_argument(
        "--max-questions", type=int, default=None,
        help="Limit number of questions (for testing)",
    )
    parser.add_argument(
        "--use-proxy", action="store_true",
        help="Use proxy for API calls",
    )
    parser.add_argument(
        "--delay", type=float, default=3.0,
        help="Delay between API calls (seconds)",
    )
    parser.add_argument(
        "--output", type=str, default=None,
        help="Custom output path for Excel file",
    )
    args = parser.parse_args()

    if not ALFRED_CLIENT_APP:
        raise SystemExit("Missing Alfred client app name. Set ALFRED_CLIENT_APP_NAME.")
    if not ALFRED_API_KEY:
        raise SystemExit("Missing Alfred API key. Set ST_CHATGPT_API_KEY (or ST_AI_BRIDGE_API_KEY/ST_API_KEY).")
    if not PERSONA_CLIENT_APP:
        raise SystemExit("Missing persona client app name. Set PERSONA_CLIENT_APP_NAME.")
    if not PERSONA_API_KEY:
        raise SystemExit("Missing persona API key. Set PERSONA_API_KEY or ST_GITHUB_ANALYZER_API_KEY.")

    questions = COMPARISON_QUESTIONS
    if args.max_questions:
        questions = questions[:args.max_questions]

    output_path = Path(args.output) if args.output else (
        OUTPUT_DIR / "comparison_alfred_vs_persona.xlsx"
    )

    print("=" * 70)
    print("  COMPARISON: Alfred (generic) vs ST GitHub Analyzer (KB)")
    print("=" * 70)
    print(f"  Questions: {len(questions)}")
    print(f"  Output: {output_path}")
    print(f"  Proxy: {'yes' if args.use_proxy else 'no'}")
    print("=" * 70)
    print()

    results = []

    for idx, q in enumerate(questions, start=1):
        question = q["question"]
        print(f"[{idx}/{len(questions)}] ({q['tags']}) {question[:70]}...")

        # Call Alfred
        print("  → Calling Alfred...")
        alfred_answer = ask_alfred(question, use_proxy=args.use_proxy)
        print(f"    Alfred: {alfred_answer[:100]}...")

        time.sleep(args.delay)

        # Call Persona
        print("  → Calling ST GitHub Analyzer...")
        persona_answer = ask_persona(question, use_proxy=args.use_proxy)
        print(f"    Persona: {persona_answer[:100]}...")

        results.append({
            "question": question,
            "category": q["category"],
            "tags": q["tags"],
            "expected_advantage": q["expected_advantage"],
            "why": q["why"],
            "alfred_answer": alfred_answer,
            "persona_answer": persona_answer,
        })

        # Save progress every 5 questions (crash-resilience: avoid losing
        # already-collected API answers if the run is interrupted).
        if idx % 5 == 0:
            export_comparison_excel(results, output_path)
            print(f"  [Progress saved: {idx}/{len(questions)}]")

        print()
        time.sleep(args.delay)

    # Final export
    export_comparison_excel(results, output_path)

    # Also save raw JSON for further analysis
    json_path = output_path.with_suffix(".json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)
    print(f"JSON saved: {json_path}")

    # Print summary
    print()
    print("=" * 70)
    print("  SUMMARY")
    print("=" * 70)
    print(f"  Total questions: {len(results)}")
    print(f"  Categories: {set(r['category'] for r in results)}")
    print(f"  Expected Persona wins: {sum(1 for r in results if r['expected_advantage'] == 'Persona')}")
    print(f"  Expected Ties: {sum(1 for r in results if r['expected_advantage'] == 'Tie')}")
    print()
    print("  Next steps:")
    print("  1. Open the Excel and fill columns 'Avantage réel', 'cite issue?', 'actionnable?'")
    print("  2. Score each answer using the grading rubric (Sheet 2)")
    print("  3. Compare total scores: Alfred vs Persona")
    print("  4. Use results in PFE report Chapter 6 (Évaluation)")
    print("=" * 70)


if __name__ == "__main__":
    main()
