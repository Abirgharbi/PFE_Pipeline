"""
Generate a per-series quality test suite (Question, Answer, Tags) covering every
STM32Cube series configured for the pipeline, and export it as an Excel workbook
with an empty Answer column ready to be filled (manually, or via
``fill_answers_from_alfred.py``).

The list of series is NOT read from datasets/*.json. It is inferred from the
pipeline configuration files:
    shared/config/config.json          (default/main series, e.g. H7)
    shared/config/config_<series>.json (one file per additional series, e.g. config_wl3.json)
(``config_all_series.json`` is intentionally excluded — it aggregates all series
and would otherwise duplicate every question set).

For each series config, 4 generic troubleshooting questions are generated from
fixed templates (DMA HardFault, package migration/breaking-change checklist,
I2C BUSY recovery, low-power audit), parameterized with the series' main repo name.

Output:
    datasets/07_delivery/st_ready/quality_test_suite_all_series.xlsx
    Sheet "Quality Test Suite" with columns:
      A = Question (width 95), B = Answer (width 70, left empty), C = Tags (width 14)
    Header row is bold/white-on-blue, data rows wrap text, all cells have thin
    borders, header row is frozen and an AutoFilter is applied.

Usage:
    python pipeline_Automation/test_suites/generate_quality_test_suite.py
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    raise SystemExit(1)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
CONFIG_DIR = PROJECT_ROOT / "shared" / "config"
OUTPUT_DIR = PROJECT_ROOT / "datasets" / "07_delivery" / "st_ready"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
XLSX_PATH = OUTPUT_DIR / "quality_test_suite_all_series.xlsx"


def _load_json(path: Path) -> dict:
    """Load a JSON config file, tolerating a UTF-8 byte-order mark (utf-8-sig)."""
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _iter_series_configs() -> list[Path]:
    """Return the ordered list of per-series config files to process.

    Includes shared/config/config.json (the default/main series) first, then
    every shared/config/config_*.json file in sorted order, except
    config_all_series.json which aggregates all series and is not a per-series
    config on its own.
    """
    paths: list[Path] = []

    default_cfg = CONFIG_DIR / "config.json"
    if default_cfg.exists():
        paths.append(default_cfg)

    for p in sorted(CONFIG_DIR.glob("config_*.json")):
        if p.name == "config_all_series.json":
            continue
        paths.append(p)

    return paths


def _find_main_repo(cfg: dict, fallback: str) -> str:
    """Pick the representative repo name for a series config.

    Prefers the first repo whose name starts with "stm32cube" (case-insensitive)
    since a config can list multiple related repos (BSP, examples, etc.) and we
    only need one canonical name to plug into the question templates. Falls
    back to the first listed repo, or to ``fallback`` if the list is empty.
    """
    repos = cfg.get("repos", [])
    if isinstance(repos, list):
        for repo in repos:
            r = str(repo).strip()
            if r.lower().startswith("stm32cube"):
                return r
        if repos:
            return str(repos[0]).strip() or fallback
    return fallback


def _series_tag_from_cfg(cfg_path: Path, cfg: dict) -> str:
    """Derive a short series tag (e.g. "H7", "WL3") used to label questions.

    For the default config.json, the tag is derived from the main repo name
    (e.g. "STM32CubeH7" -> "H7"), defaulting to "H7" if it can't be parsed.
    For config_<series>.json files, the tag is the uppercased <series> part
    of the filename.
    """
    if cfg_path.name == "config.json":
        main = _find_main_repo(cfg, "STM32CubeH7")
        if main.lower().startswith("stm32cube"):
            return main[len("STM32Cube") :].upper() or "H7"
        return "H7"

    stem = cfg_path.stem  # config_wl3
    tag = stem.replace("config_", "", 1).upper()
    return tag


def _build_questions_for_series(tag: str, main_repo: str) -> list[dict[str, str]]:
    """Instantiate the fixed question templates for one series.

    Returns a list of dicts with empty "answer" (to be filled later) and the
    series "tags" so results can be grouped/filtered in the Excel sheet.
    """
    # 4 robust questions per series keeps runtime reasonable while covering common failure modes.
    templates = [
        (
            "My {main_repo} project crashes with HardFault when enabling DMA in interrupt mode. "
            "Polling mode works. What should I verify first in clock, memory alignment, and interrupt configuration?"
        ),
        (
            "After updating to a newer {main_repo} package release, a previously working peripheral init now fails. "
            "What is the best migration/debug checklist to identify breaking changes quickly?"
        ),
        (
            "On {main_repo}, I2C sometimes gets stuck BUSY after a communication error and the bus never recovers. "
            "What is a safe software recovery sequence without full MCU reset?"
        ),
        (
            "I need lower power consumption on {main_repo} while preserving wake-up reliability. "
            "Which low-power configuration points should be audited first (clocks, GPIO state, peripherals, wake-up sources)?"
        ),
    ]

    return [
        {
            "question": t.format(tag=tag, main_repo=main_repo),
            "answer": "",
            "tags": tag,
        }
        for t in templates
    ]


def build_questions() -> list[dict[str, str]]:
    """Build the full question list by iterating over every series config.

    For each discovered config file, resolves its series tag and main repo
    name, then generates that series' 4 template questions and appends them
    to a single flat list (used to populate the ``QUESTIONS`` module constant).
    """
    all_questions: list[dict[str, str]] = []

    for cfg_path in _iter_series_configs():
        cfg = _load_json(cfg_path)
        tag = _series_tag_from_cfg(cfg_path, cfg)
        main_repo = _find_main_repo(cfg, f"STM32Cube{tag}")
        all_questions.extend(_build_questions_for_series(tag=tag, main_repo=main_repo))

    return all_questions


# Built once at import time so other scripts (e.g. fill_answers_from_alfred.py)
# can simply do `from generate_quality_test_suite import QUESTIONS`.
QUESTIONS = build_questions()


def backup_if_exists(path: Path) -> Path | None:
    """Create a side-by-side .bak copy when the target output already exists."""
    if not path.exists():
        return None
    backup_path = path.with_suffix(path.suffix + ".bak")
    shutil.copy2(path, backup_path)
    return backup_path


def main() -> None:
    """Build the question list and write it to a styled .xlsx workbook.

    The Answer column is left blank; it is meant to be filled in afterwards,
    either manually or by ``fill_answers_from_alfred.py``.
    """
    backup_path = backup_if_exists(XLSX_PATH)
    if backup_path is not None:
        print(f"Existing file backed up to: {backup_path}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Quality Test Suite"

    header_fill = PatternFill(start_color="1A2382", end_color="1A2382", fill_type="solid")
    header_font = Font(bold=True, size=11, color="FFFFFF")
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["Question", "Answer", "Tags"]
    col_widths = [95, 70, 14]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[chr(64 + col_idx)].width = width

    for row_idx, q in enumerate(QUESTIONS, start=2):
        row_data = [q["question"], q["answer"], q["tags"]]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align
            cell.border = thin_border
            cell.font = Font(size=10)

    # Keep the header row visible and enable column filtering/sorting for graders.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(QUESTIONS) + 1}"

    wb.save(XLSX_PATH)

    tags = sorted({q["tags"] for q in QUESTIONS})
    print(f"Generated: {XLSX_PATH}")
    print(f"Total questions: {len(QUESTIONS)}")
    print(f"Series covered ({len(tags)}): {', '.join(tags)}")


if __name__ == "__main__":
    main()
