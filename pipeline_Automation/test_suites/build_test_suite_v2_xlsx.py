"""
Convert the v2 CSV test suite (Question;Answer;Tags) into a formatted Excel workbook.

This is the v2 variant of the CSV-to-XLSX converter used for manual QA grading of
the knowledge base. It does not read any pipeline JSON dataset; it only reformats
an already-generated CSV file. Unlike the v1 builder, this script runs as a plain
module-level script (no ``main()``/``if __name__`` guard) and does not strip quote
characters from cell values.

Input:
    datasets/07_delivery/st_ready/all_series_test_suite_50q_v2.csv
    (semicolon-delimited: Question;Answer;Tags)

Output:
    datasets/07_delivery/st_ready/all_series_test_suite_50q_v2.xlsx
    Sheet "Test Suite 50Q v2" with columns:
      A = Question (width 90), B = Answer (width 100), C = Tags (width 20)
    Header row is bold/white-on-blue and centered, data rows wrap text, all cells
    have thin borders, and the header row is frozen.

Usage:
    python pipeline_Automation/test_suites/build_test_suite_v2_xlsx.py
"""
import csv
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    raise SystemExit(1)

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "datasets" / "07_delivery" / "st_ready" / "all_series_test_suite_50q_v2.csv"
XLSX_PATH = CSV_PATH.with_suffix(".xlsx")

wb = Workbook()
ws = wb.active
ws.title = "Test Suite 50Q v2"

# Styles
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

with open(CSV_PATH, encoding="utf-8") as f:
    reader = csv.reader(f, delimiter=";")
    for row_idx, row in enumerate(reader, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if row_idx == 1:
                # Header row: bold white-on-blue, centered so it reads like a
                # table title rather than a data cell.
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            else:
                # Data rows: wrap long question/answer text and align to the
                # top so multi-line cells stay readable.
                cell.alignment = Alignment(wrap_text=True, vertical="top")

# Column widths sized for long free-text Question/Answer content plus a short Tags column.
ws.column_dimensions["A"].width = 90   # Question
ws.column_dimensions["B"].width = 100  # Answer
ws.column_dimensions["C"].width = 20   # Tags

# Keep the header row visible while scrolling through many questions.
ws.freeze_panes = "A2"

wb.save(XLSX_PATH)
print(f"Created: {XLSX_PATH}")
print(f"Rows: {ws.max_row - 1} questions")
