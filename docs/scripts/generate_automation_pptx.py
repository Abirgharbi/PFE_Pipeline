from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def add_title_slide(prs: Presentation, title: str, subtitle: str, notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    slide.notes_slide.notes_text_frame.text = notes


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    for idx, line in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(22)

    slide.notes_slide.notes_text_frame.text = notes


def add_two_column_slide(
    prs: Presentation,
    title: str,
    left_title: str,
    left_lines: list[str],
    right_title: str,
    right_lines: list[str],
    notes: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    left = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(5.7), Inches(4.8))
    left_tf = left.text_frame
    left_tf.text = left_title
    left_tf.paragraphs[0].font.bold = True
    left_tf.paragraphs[0].font.size = Pt(24)
    for line in left_lines:
        p = left_tf.add_paragraph()
        p.text = f"- {line}"
        p.level = 0
        p.font.size = Pt(20)

    right = slide.shapes.add_textbox(Inches(6.2), Inches(1.6), Inches(6.6), Inches(4.8))
    right_tf = right.text_frame
    right_tf.text = right_title
    right_tf.paragraphs[0].font.bold = True
    right_tf.paragraphs[0].font.size = Pt(24)
    for line in right_lines:
        p = right_tf.add_paragraph()
        p.text = f"- {line}"
        p.level = 0
        p.font.size = Pt(20)

    slide.notes_slide.notes_text_frame.text = notes


def build_presentation(output_path: Path) -> None:
    prs = Presentation()

    add_title_slide(
        prs,
        title="GitHub Actions Automation for STM32Cube KB Update",
        subtitle="How auto_update_kb.yml orchestrates preprocessing and upload to KB #793",
        notes=(
            "This presentation explains the automation architecture, runtime parameters, "
            "and operating runbook."
        ),
    )

    add_bullets_slide(
        prs,
        title="End-to-End Flow",
        bullets=[
            "Trigger: schedule (1st/15th) or manual run",
            "Prepare job normalizes inputs and builds matrix",
            "Run-series job executes on self-hosted Windows runner",
            "PowerShell orchestrator runs full series pipeline",
            "Artifacts and execution summary are always published",
        ],
        notes=(
            "The workflow is split into prepare and execution jobs to isolate validation "
            "from heavy processing."
        ),
    )

    add_bullets_slide(
        prs,
        title="YML Structure",
        bullets=[
            "on: schedule + workflow_dispatch inputs",
            "permissions: contents read",
            "concurrency: prevent overlap on same ref",
            "env: default series, mode, policy, python version",
            "jobs: prepare and run-series",
        ],
        notes=(
            "The file is intentionally structured from trigger and governance to runtime jobs."
        ),
    )

    add_bullets_slide(
        prs,
        title="Manual Inputs and Effects",
        bullets=[
            "series: CSV list of target series",
            "mode: Full, Prepare, Upload",
            "existing_datasource_mode: Replace or Add",
            "skip_drivers: skip subrepo uploads",
            "skip_schema_validation: bypass schema gate",
            "continue_on_workflow_error: continue inside workflow",
            "placeholder_policy: Fail, Warn, Off",
            "max_parallel: series concurrency",
        ],
        notes=(
            "Each input maps directly to the PowerShell script arguments for transparent control."
        ),
    )

    add_bullets_slide(
        prs,
        title="Per-Series Execution",
        bullets=[
            "checkout repository",
            "enable Windows long paths",
            "setup Python 3.12 and install requirements",
            "run Run_Single_Series_Full_Pipeline_And_Upload.ps1",
            "upload artifacts and publish run summary",
        ],
        notes=(
            "Artifact and summary steps use always() logic in the workflow to keep traceability "
            "even when one step fails."
        ),
    )

    add_two_column_slide(
        prs,
        title="Reliability and Governance",
        left_title="Reliability",
        left_lines=[
            "fail-fast disabled for matrix isolation",
            "timeout per series job",
            "placeholder policy gate",
            "schema validation gate",
        ],
        right_title="Traceability",
        right_lines=[
            "per-series artifact archive",
            "GitHub step summary",
            "datasource IDs persisted by script",
            "manual rerun with mode Upload",
        ],
        notes=(
            "This design supports maintenance operations: isolated failures, clear diagnostics, "
            "and resumable uploads."
        ),
    )

    add_bullets_slide(
        prs,
        title="Operator Runbook",
        bullets=[
            "ensure self-hosted runner is online",
            "verify secrets ST_CHATGPT_API_KEY and ST_REMOTE_USER",
            "run smoke test: one series, mode Prepare",
            "promote to Full mode on selected series",
            "verify job status, artifacts, and summary",
        ],
        notes=(
            "Recommended first run: series G4, mode Prepare, max_parallel 1."
        ),
    )

    add_bullets_slide(
        prs,
        title="Reference Files",
        bullets=[
            ".github/workflows/auto_update_kb.yml",
            "pipeline_Automation/workflow/Run_Single_Series_Full_Pipeline_And_Upload.ps1",
            "docs/automation_kb_update.md",
            "docs/automation_kb_update_presentation.md",
        ],
        notes=(
            "These files are the minimum set to understand, run, and maintain the automation."
        ),
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate automation PowerPoint presentation.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/automation_kb_update_presentation.pptx"),
        help="Output PPTX path.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_presentation(args.output)
    print(f"PPTX generated: {args.output}")


if __name__ == "__main__":
    main()
