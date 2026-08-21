from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def add_title_and_subtitle(slide, title: str, subtitle: str) -> None:
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str], notes: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    for idx, line in enumerate(bullets):
        p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def normalize(text: str) -> str:
    return " ".join((text or "").lower().replace("-", " ").split())


def shape_text(shape) -> str:
    if hasattr(shape, "text"):
        return shape.text or ""
    return ""


def slide_title(slide) -> str:
    if slide.shapes.title is not None and hasattr(slide.shapes.title, "text"):
        return slide.shapes.title.text or ""
    for sh in slide.shapes:
        txt = shape_text(sh).strip()
        if txt:
            return txt.splitlines()[0]
    return ""


def set_slide_title(slide, title: str) -> None:
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
        return

    # Some layouts are fully blank and do not expose a title placeholder.
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True


def slide_full_text(slide) -> str:
    return "\n".join(shape_text(sh) for sh in slide.shapes if shape_text(sh).strip())


def find_slide_by_keywords(prs: Presentation, any_keywords: list[str]) -> int | None:
    for idx, slide in enumerate(prs.slides):
        hay = normalize(slide_title(slide) + "\n" + slide_full_text(slide))
        if all(normalize(k) in hay for k in any_keywords):
            return idx
    return None


def find_best_automation_slide(prs: Presentation) -> int | None:
    # Priority 1: explicit automation title.
    for idx, slide in enumerate(prs.slides):
        t = normalize(slide_title(slide))
        if "automation" in t:
            return idx

    # Priority 2: slide containing empty-ish automation placeholder text.
    for idx, slide in enumerate(prs.slides):
        full = normalize(slide_full_text(slide))
        if "automation" in full and len(full) < 120:
            return idx
    return None


def remove_shape_if_matches(shape, predicates: list[str]) -> bool:
    txt = normalize(shape_text(shape))
    if not txt:
        return False
    if any(normalize(p) in txt for p in predicates):
        sp = shape._element
        sp.getparent().remove(sp)
        return True
    return False


def add_reference_box(
    slide,
    title: str,
    lines: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    fill_rgb: tuple[int, int, int],
    line_rgb: tuple[int, int, int],
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape = box
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
    shape.line.color.rgb = RGBColor(*line_rgb)
    shape.line.width = Pt(1.5)

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.bold = True
    p0.font.size = Pt(11)

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(10)
        p.level = 0


def remove_formula_boxes(slide) -> None:
    tokens = [
        "reference formule alfred coherence",
        "reference score strategie kb + qi",
        "formula used trust",
        "formula used quality gain",
    ]
    for sh in list(slide.shapes):
        remove_shape_if_matches(sh, tokens)


def fill_automation_slide(slide, diagrams_root: Path) -> None:
    set_slide_title(slide, "Automation serie complete - schema operationnel")

    # Cleanup old temporary boxes if rerun.
    cleanup_tokens = [
        "mode full = prepare + upload",
        "slide ajoutee automatiquement",
        "ajout schema automation serie complete",
    ]
    for sh in list(slide.shapes):
        remove_shape_if_matches(sh, cleanup_tokens)

    # Prefer the new exported image if available; otherwise fallback to text summary.
    candidate_images = [
        diagrams_root / "exported" / "07_automation_series_pipeline.png",
        diagrams_root / "exported" / "06_cicd_automation.png",
    ]
    image_path = next((p for p in candidate_images if p.exists()), None)

    if image_path is not None:
        slide.shapes.add_picture(str(image_path), Inches(0.35), Inches(1.25), width=Inches(12.6))
    else:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.3), Inches(4.9))
        tf = box.text_frame
        tf.clear()
        lines = [
            "Mode Full = Prepare + Upload; Mode Prepare = artefacts locaux; Mode Upload = publication KB.",
            "Etape 1-2: sync repo/submodules puis preprocessing complet avec Alfred PDF pendant le run.",
            "Etape 3-4: exports ST-ready parent+drivers puis enrichissement Alfred des issues images.",
            "Etape 5-6: patch PDF descriptions puis validation schema + gate placeholders.",
            "Etape 7-8: creation/add datasources KB #793 puis upload drivers/subrepos avec split JSON.",
            "Sortie: upload_datasource_ids_<serie>.json + sync vers shared/config/config_all_series.json.",
        ]
        for idx, line in enumerate(lines):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(18)
            p.level = 0

    note = (
        "Diapo automation finalisee en place. Le schema PNG 07_automation_series_pipeline "
        "est prioritaire; sinon le resume texte est insere."
    )
    slide.notes_slide.notes_text_frame.text = note


def add_formula_block(tf, text: str, size: int = 16, bold: bool = False) -> None:
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold


def inject_formula_references(prs: Presentation) -> None:
    # Target 1: "Alfred Vision - Contextual Trust" slide.
    alfred_idx = find_slide_by_keywords(prs, ["alfred vision", "contextual trust"])
    if alfred_idx is not None:
        slide = prs.slides[alfred_idx]
        remove_formula_boxes(slide)
        add_reference_box(
            slide,
            title="Formula used (trust)",
            lines=[
                "Trust(%) = (Evaluated - LowCoherence) / Evaluated * 100",
                "Issue: (200 - 3) / 200 = 98.5%",
                "PDF: (242 - 141) / 242 = 41.7%",
            ],
            left=8.9,
            top=1.22,
            width=3.8,
            height=1.05,
            fill_rgb=(238, 251, 251),
            line_rgb=(0, 137, 123),
        )

    # Target 2: "Comparative Study - Generic LLM vs Persona KB" slide.
    persona_idx = find_slide_by_keywords(prs, ["comparative study", "generic llm", "persona kb"])
    if persona_idx is not None:
        slide = prs.slides[persona_idx]
        remove_formula_boxes(slide)
        add_reference_box(
            slide,
            title="Formula used (quality gain)",
            lines=[
                "Gain(%) = (ScorePersona - ScoreLLM) / ScoreLLM * 100",
                "= (8.1 - 5.8) / 5.8 * 100 = 39.7% ~ 39%",
                "Score/answer = mean(question scores on /10)",
            ],
            left=8.9,
            top=1.18,
            width=3.8,
            height=1.05,
            fill_rgb=(239, 247, 255),
            line_rgb=(25, 118, 210),
        )


def remove_extra_generated_slides(prs: Presentation) -> None:
    # Remove slides previously injected by the old script.
    to_remove = {
        "plan mise a jour automation + evaluation",
        "automation serie complete schema operationnel",
        "evaluation alfred et comparaison persona formules",
    }

    indices = []
    for i, slide in enumerate(prs.slides):
        t = normalize(slide_title(slide))
        if t in to_remove:
            indices.append(i)

    # Delete from back to front.
    for idx in sorted(indices, reverse=True):
        slide_id = prs.slides._sldIdLst[idx]
        prs.slides._sldIdLst.remove(slide_id)


def build_updated_presentation(input_pptx: Path, output_pptx: Path, diagrams_root: Path) -> None:
    prs = Presentation(str(input_pptx))

    remove_extra_generated_slides(prs)

    automation_idx = find_best_automation_slide(prs)
    if automation_idx is not None:
        fill_automation_slide(prs.slides[automation_idx], diagrams_root=diagrams_root)

    inject_formula_references(prs)

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Update STM32Cube_worshopF1.pptx with automation and evaluation content."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("docs/STM32Cube_worshopF1.pptx"),
        help="Input PPTX to update.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/STM32Cube_worshopF1.pptx"),
        help="Output PPTX path (same as input to update in-place).",
    )
    parser.add_argument(
        "--diagrams-root",
        type=Path,
        default=Path("docs/diagrams/presentation"),
        help="Root folder where drawio exported PNG diagrams are located.",
    )
    return parser.parse_args()


def resolve_existing_input(path_arg: Path) -> Path:
    if path_arg.is_absolute() and path_arg.exists():
        return path_arg

    script_dir = Path(__file__).resolve().parent
    repo_root = script_dir.parents[1]
    docs_dir = repo_root / "docs"

    candidates = [
        Path.cwd() / path_arg,
        script_dir / path_arg,
        docs_dir / path_arg,
        repo_root / path_arg,
    ]

    for candidate in candidates:
        if candidate.exists():
            return candidate

    return path_arg


def resolve_output_path(path_arg: Path) -> Path:
    if path_arg.is_absolute():
        return path_arg

    script_dir = Path(__file__).resolve().parent
    repo_root = script_dir.parents[1]
    docs_dir = repo_root / "docs"

    # If user provides only filename, default output under docs/.
    if len(path_arg.parts) == 1:
        return docs_dir / path_arg

    # Otherwise keep relative to current working directory.
    return (Path.cwd() / path_arg).resolve()


def main() -> None:
    args = parse_args()
    resolved_input = resolve_existing_input(args.input)
    resolved_output = resolve_output_path(args.output)

    if not resolved_input.exists():
        raise FileNotFoundError(
            "Input PPTX not found. Tried path: "
            f"{args.input}. You can pass an absolute path or one under docs/."
        )

    build_updated_presentation(
        input_pptx=resolved_input,
        output_pptx=resolved_output,
        diagrams_root=args.diagrams_root,
    )
    print(f"Updated PPTX written to: {resolved_output}")


if __name__ == "__main__":
    main()
