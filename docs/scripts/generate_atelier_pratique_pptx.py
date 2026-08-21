from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


def add_title_slide(prs: Presentation, title: str, subtitle: str, notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    slide.notes_slide.notes_text_frame.text = notes


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    for idx, line in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(21)

    slide.notes_slide.notes_text_frame.text = notes


def build_presentation(output_path: Path) -> None:
    prs = Presentation()

    add_title_slide(
        prs,
        title="Atelier Pratique - Pipeline STM32Cube",
        subtitle="Ingestion a Upload KB #793 avec gouvernance Add/Replace",
        notes=(
            "Version atelier mise a jour apres les modifications de l'automatisation "
            "(runner-check, fallback IDs, sync config globale)."
        ),
    )

    add_bullets_slide(
        prs,
        title="Objectif Atelier",
        bullets=[
            "Comprendre le flux de bout en bout: Ingestion -> Delivery -> Upload",
            "Executer une serie en mode Full, Prepare ou Upload",
            "Garantir un upload robuste et tracable vers la KB #793",
            "Maintenir les IDs datasource coherents dans les snapshots et la config globale",
        ],
        notes="Cette slide introduit les outcomes attendus de l'atelier pratique.",
    )

    add_bullets_slide(
        prs,
        title="Architecture CI/CD Actuelle",
        bullets=[
            "Workflow: .github/workflows/auto_update_kb.yml",
            "prepare: normalise les inputs et construit la matrice series",
            "runner-check: verifie la disponibilite du self-hosted runner",
            "run-series: execute le script PowerShell serie par serie",
            "runner-offline-notice: publie Upload Deferred si runner offline",
        ],
        notes="Insister que l'upload n'est lance que si un runner self-hosted est online.",
    )

    add_bullets_slide(
        prs,
        title="Commande Atelier Recommandee",
        bullets=[
            "Run_Single_Series_Full_Pipeline_And_Upload.ps1 centralise prepare + upload",
            "Exemple: -Series WB -Mode Full -ExistingDatasourceMode Replace",
            "Mode Prepare: genere et valide sans upload",
            "Mode Upload: rejoue uniquement la phase upload",
            "placeholder_policy Fail protege contre les placeholders non resolus",
        ],
        notes="Presenter Full pour run complet, Upload pour reprise rapide apres incident.",
    )

    add_bullets_slide(
        prs,
        title="Add vs Replace",
        bullets=[
            "Add: ajoute de nouveaux fichiers dans une datasource existante",
            "Replace: nettoie puis republie pour garder la KB a jour sans doublons",
            "Replace active un controle strict fail-on-existing-files",
            "Le choix est passe de bout en bout (workflow -> script -> uploader)",
        ],
        notes="Recommander Replace pour production, Add pour cas incrementaux cibles.",
    )

    add_bullets_slide(
        prs,
        title="Robustesse sur IDs Datasource",
        bullets=[
            "Si un ID sauvegarde est perime: recreation automatique de la datasource",
            "Fallback actif en mode Replace et en mode Add",
            "Snapshot incremental: issues, files, diagnostic, resolver",
            "last_successful_stage permet une reprise claire apres echec",
        ],
        notes=(
            "Cette evolution corrige le cas 'KbDataSourceEntity not found' observe sur WB."
        ),
    )

    add_bullets_slide(
        prs,
        title="Sync Config Globale",
        bullets=[
            "Les IDs uploades sont ecrits dans upload_datasource_ids_<slug>.json",
            "Le script synchronise aussi shared/config/config_all_series.json",
            "Mise a jour automatique des cles issues/files/diagnostic/resolver",
            "Exemple WB: IDs remplis automatiquement apres upload",
        ],
        notes=(
            "Evite les ecarts entre snapshots de serie et kb_datasource_ids globale."
        ),
    )

    add_bullets_slide(
        prs,
        title="Validation et Gates",
        bullets=[
            "Schema validation avant upload (sauf skip explicite)",
            "Placeholder gate (Fail/Warn/Off) pour images/PDF Alfred",
            "fail-fast=false sur matrix pour isoler les erreurs par serie",
            "Artifacts et step summaries conserves pour audit",
        ],
        notes="Montrer comment la qualite est protegee avant publication dans la KB.",
    )

    add_bullets_slide(
        prs,
        title="Runbook Operateur",
        bullets=[
            "1) Lancer un pilote en mode Prepare",
            "2) Verifier summaries et validate_schemas",
            "3) Lancer Full en Replace pour mise a jour propre",
            "4) Utiliser Upload pour reprise rapide",
            "5) Verifier kb_datasource_ids dans config_all_series.json",
        ],
        notes="Runbook a suivre pendant l'atelier et en exploitation courante.",
    )

    add_bullets_slide(
        prs,
        title="Conclusion",
        bullets=[
            "Pipeline outille pour production: traceable, resumable, governable",
            "Ajouts recents: fallback IDs + sync config globale + check runner",
            "Prochaine etape infra: runner dedie always-on (VM/serveur)",
        ],
        notes="Clore sur la maturite operationnelle et le plan d'industrialisation.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate Atelier Pratique PPTX.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/Atelier_Pratique_Pipeline_STM32Cube.pptx"),
        help="Output PPTX path.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_presentation(args.output)
    print(f"PPTX generated: {args.output}")


if __name__ == "__main__":
    main()
