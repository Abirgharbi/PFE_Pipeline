from pptx import Presentation
p = Presentation(r"c:\Users\gharbiab\OneDrive - STMicroelectronics\Desktop\PFE_Chatbot_STM32Cube\docs\Atelier_Pratique_Pipeline_STM32Cube.pptx")
for idx, s in enumerate(p.slides, 1):
    title = s.shapes[0].text if len(s.shapes) and hasattr(s.shapes[0], 'text') else ''
    if 'Checklist' in title:
        print('slide', idx, title)
        for i, sh in enumerate(s.shapes):
            if hasattr(sh, 'text') and sh.text:
                print('shape', i, sh.text[:500].replace('\n',' | '))
